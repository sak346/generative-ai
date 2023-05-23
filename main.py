import os
import re
import tkinter as tk
import webbrowser
from tkinter import filedialog

import PyPDF2
from langchain import OpenAI
from langchain.chains.summarize import load_summarize_chain
from langchain.docstore.document import Document
from langchain.prompts import PromptTemplate
from langchain.text_splitter import CharacterTextSplitter
import openpyxl

root = tk.Tk()
root.withdraw()

os.environ["OPENAI_API_KEY"] = "sk-4hcv1t7csGC4ZnfKk0vpT3BlbkFJvebKDubyOuRwDGWAS7RS"


def first_layout(bullet_points):
    html_code = "<html><head><style>"
    html_code += ".box { width: 300px; height:100px;  margin: 20px; padding: 30px 20px 20px 20px; text-align: center;overflow: hidden;" \
                 " border:solid; box-shadow:3px; border-radius:10px;}"
    html_code += ".title{ text-align: center;font-size:40px; }"
    html_code += ".outer-box{ padding:20px ; width: 100% ; align:centre;}"
    html_code += "</style></head><body style="'background-color:#F8F0E3;' ">"
    html_code += "<div class='title'>Action Items</div>"
    html_code += "<div class='outer-box'>"

    for i, item in enumerate(bullet_points):
        html_code += f"<div class='box'>{item}</div>"

    html_code += "</div>"
    html_code += "</body></html>"

    file_name = 'action_items_layout_1.html'
    with open(file_name, "w") as file:
        file.write(html_code)
    html_file = file_name
    webbrowser.open(os.path.abspath(html_file))
    input("Press Enter to continue")
    document_file = 'output_document_layout_1.html'
    with open(document_file, "w") as file:
        file.write(html_code)


def second_layout(bullet_points):
    topic = " "

    html_code = " <html> <head> <style>"
    html_code += " body {  justify-content: center;  align-items: center; height: 100vh; background-color: #f2f2f2;}"
    html_code += ".title{ text-align: center; font-size:40px;}"
    html_code += " .container { text-align: center;   padding: 20px; margin-top: 40px; }"
    html_code += " .box { display: inline-block;  padding: 10px;  text-align: left; background-color: #f5f5f5; }"
    html_code += " .action-items { list-style-type: none;padding: 5px; margin-top: 10px; }"
    html_code += ".action-items li { margin-bottom: 10px; font-size:35px;}"
    html_code += "  </style> </head> <body>"
    html_code += " <div class='title'>Action Items</div>"
    html_code += f'''<div class='topic' style=" color: red; text-align: center; ">{topic}</div>'''
    html_code += " <div class='container'> <div class='box'>"
    html_code += "<ul class='action-items'>"
    html_code += f'''{''.join(f"<li>{item}</li>" for item in bullet_points)}'''
    html_code += "</ul> </div> </div> </body> </html>"

    file_name = 'action_items_layout_2.html'

    with open(file_name, "w") as file:
        file.write(html_code)
    html_file = 'action_items_layout_2.html'
    webbrowser.open(os.path.abspath(html_file))

    input("Press Enter to continue")
    document_file = 'output_document_layout_2.html'
    with open(document_file, "w") as file:
        file.write(html_code)


def third_layout(bullet_points):
    html_code = "<html><head><style>"
    html_code += ".box { width: 300px; margin: 20px; padding: 20px; text-align: center;overflow: hidden;" \
                 " border:solid; box-shadow:3px; border-radius:10px;}"
    html_code += ".title{ text-align: center;font-size:40px; }"
    html_code += ".outer-box{ display:flex; width: 100% ; flex-wrap: wrap;}"
    html_code += "</style></head><body style="'background-color:#F8F0E3;' ">"
    html_code += "<div class='title'>Action Items</div>"
    html_code += "<div class='outer-box'>"

    for i, item in enumerate(bullet_points):
        html_code += f"<div class='box'>{item}</div>"

    html_code += "</div>"
    html_code += "</body></html>"

    file_name = 'action_items_layout_2.html'

    with open(file_name, "w") as file:
        file.write(html_code)
    html_file = file_name
    webbrowser.open(os.path.abspath(html_file))
    input("Press Enter to continue")
    document_file = 'output_document_layout_2.html'
    with open(document_file, "w") as file:
        file.write(html_code)


def is_pdf_file(file_path):
    _, file_ext = os.path.splitext(file_path)
    return file_ext.lower() == '.pdf'


def is_excel_file(file_path):
    _, file_ext = os.path.splitext(file_path)
    return file_ext.lower() == '.xlsx'


def cleanTxtFile(text_file):
    with open(text_file, 'r', encoding='utf-8') as file:
        lines = file.readlines()

    cleaned_lines = [line.strip() for line in lines if line.strip()]

    with open(text_file, 'w', encoding='utf-8') as file:
        file.write('\n'.join(cleaned_lines))


def convert_pdf_to_txt(pdf_file, text_file):
    with open(pdf_file, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()

    with open(text_file, 'w', encoding='utf-8') as file:
        file.write(text)

    cleanTxtFile(text_file)

    print(f"PDF converted to {text_file}")


def convert_xlsx_to_txt(xls_file, text_file):
    workbook = openpyxl.load_workbook(xls_file)
    sheet = workbook.active
    with open(text_file, 'w') as file:
        for row in sheet.iter_rows(values_only=True):
            line = '\t'.join(str(cell) for cell in row)
            file.write(line + '\n')

    cleanTxtFile(text_file)
    print(f"EXCEL converted to {text_file}")

def create_file(file_name):
    with open(file_name, 'w') as file:
        pass

    print(f"File '{file_name}' created.")


if __name__ == "__main__":
    while True:

        choice = input("Want to select document for action point extraction Yes/No?:")

        if choice.lower() == "yes":
            llm = OpenAI(temperature=0)
            text_splitter = CharacterTextSplitter(chunk_size=2000, chunk_overlap=200)
            file_path = filedialog.askopenfilename()

            if file_path:
                path = file_path
            else:
                print("No file selected.")
                break

            new_txt_file = ''
            txt_file = ''

            if is_pdf_file(path) or is_excel_file(path):
                new_file_name = '.txt'
                print(new_file_name)
                create_file(new_file_name)
                txt_file = new_file_name

                if is_pdf_file(path):
                    convert_pdf_to_txt(path, txt_file)

                elif is_excel_file(path):
                    convert_xlsx_to_txt(path, txt_file)

            if txt_file == '':
                new_txt_file = path
            else:
                new_txt_file = txt_file

            with open(new_txt_file) as f:
                document = f.read()
            texts = text_splitter.split_text(document)

            docs = [Document(page_content=t) for t in texts]

            prompt_template = """Extract action items from the document:

            {text}

            Action Items Of the document:"""

            PROMPT = PromptTemplate(template=prompt_template, input_variables=["text"])
            chain = load_summarize_chain(llm, chain_type="stuff", prompt=PROMPT)
            action_items = []

            for doc in docs:
                paragraph = chain.run([doc])
                action_items.append(paragraph)

            all_action_items = ' '.join(action_items)
            pattern = r'\d+\.(.*?)\n'
            matches = re.findall(pattern, all_action_items, re.DOTALL)
            action_sentences = [match.strip() for match in matches]

            bullet_points = action_sentences

            layout_choice = input("Select a layout (1 or 2 or 3): ")

            if layout_choice == "1":
                print("First Layout:")
                first_layout(bullet_points)

            elif layout_choice == "2":
                print("Second Layout:")
                second_layout(bullet_points)

            elif layout_choice == "3":
                print("Third Layout:")
                third_layout(bullet_points)
            else:
                print("Invalid choice. Please select either 1 , 2 or 3.")

        else:
            break
